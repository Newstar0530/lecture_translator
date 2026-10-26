"""
即時上課翻譯＋筆記：核心邏輯（不含介面）

流程：
    麥克風 / 音訊檔
      → Voxtral Realtime（語音轉文字，串流）
      → 斷句（SentenceBuffer）
      → Mistral LLM 翻譯成繁體中文
      → 顯示在 Streamlit，並即時存檔
    下課後：
      → 把整堂逐字稿分段整理，產生 Markdown 筆記
"""
from __future__ import annotations

import asyncio
import csv
import hashlib
import json
import queue as thread_queue
import os
import re
import shutil
import subprocess
import sys
import threading
import time
import wave
from collections import deque
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import AsyncIterator, Callable, Optional

import numpy as np

SAMPLE_RATE = 16000          # Voxtral Realtime 需要 16 kHz、單聲道、16-bit PCM
CHUNK_MS = 100               # 每次送出 100 毫秒的音訊
CHUNK_SAMPLES = SAMPLE_RATE * CHUNK_MS // 1000
REALTIME_PRICE_PER_MIN = 0.006   # 美元／分鐘（2026 年官方公告價格，可能變動）

DEFAULT_STT_MODEL = "voxtral-mini-transcribe-realtime-2602"
DEFAULT_TRANSLATE_MODEL = "mistral-small-2603"   # 固定版本（＝ 2026-10 的 mistral-small-latest），研究期間才不會被自動換掉
DEFAULT_NOTES_MODEL = "mistral-medium-latest"
DEFAULT_SLIDES_MODEL = "mistral-large-latest"   # 讀簡報只跑一次，用最強的模型


# ---------------------------------------------------------------------------
# 1. 斷句：把串流進來的零碎文字，組成適合翻譯的句子
# ---------------------------------------------------------------------------
_SENTENCE_END = re.compile(r"[.!?…。！？](?:[\"'”’)\]]*)\s")
# 句點前面是這些縮寫時不算句尾（英文＋法文，比對時不分大小寫）
_ABBREVIATIONS = {
    "mr", "mrs", "ms", "dr", "prof", "st", "vs", "etc", "e.g", "i.e", "cf", "fig", "no",
    "vol", "p", "pp", "ch", "approx", "jr", "sr", "inc", "ltd", "co", "u.s", "u.k",
    "m", "mme", "mlle", "ex", "env",
}


def _is_abbreviation(text: str, dot_pos: int) -> bool:
    """text[dot_pos] 是句點時，判斷它是不是縮寫（Dr.、e.g.）或人名縮寫（J. Smith）。"""
    if text[dot_pos] != ".":
        return False
    m = re.search(r"([A-Za-zÀ-ÿ.]+)$", text[:dot_pos])
    if not m:
        return False
    word = m.group(1).lstrip(".")
    return word.lower() in _ABBREVIATIONS or (len(word) == 1 and word.isupper())


class SentenceBuffer:
    """
    規則：
    - 遇到句尾標點（. ! ? …）而且後面接了空白，而且累積長度 >= min_chars → 送出
      （Dr.、e.g.、J. Smith 這類縮寫的句點不算句尾）
    - 沒有句號但已經累積 comma_seconds 秒或 comma_chars 個字元 → 在最後一個逗號處切開送出
      （語音辨識有時整段只標逗號，不這樣做的話要等老師停下來才會翻譯）
    - 累積超過 max_chars 還沒有句尾 → 在最後一個逗號或空白處切開送出
    - 一段時間沒有新文字（idle_seconds）→ 把剩下的送出（由外部呼叫 flush_if_idle）
    min_chars 可以把太短的句子合併，減少翻譯請求次數。
    """

    def __init__(self, min_chars: int = 80, max_chars: int = 320, idle_seconds: float = 3.0,
                 comma_seconds: float = 6.0, comma_chars: int = 160):
        self.min_chars = min_chars
        self.max_chars = max_chars
        self.idle_seconds = idle_seconds
        self.comma_seconds = comma_seconds
        self.comma_chars = comma_chars
        self.buf = ""
        self.last_update = time.monotonic()
        self.buf_started = self.last_update      # 目前這段開始累積的時間

    def add(self, text: str) -> list[str]:
        now = time.monotonic()
        if not self.buf.strip():
            self.buf_started = now
        self.buf += text
        self.last_update = now
        out: list[str] = []
        while True:
            cut = self._find_cut(now)
            if cut is None:
                break
            piece, self.buf = self.buf[:cut].strip(), self.buf[cut:].lstrip()
            self.buf_started = now
            if piece:
                out.append(piece)
        return out

    def _find_cut(self, now: float) -> Optional[int]:
        # 找「長度已經夠」之後的第一個句尾
        for m in _SENTENCE_END.finditer(self.buf):
            if m.end() >= self.min_chars and not _is_abbreviation(self.buf, m.start()):
                return m.end()
        # 沒有句號、但已經等太久或太長：在最後一個逗號處切
        if now - self.buf_started >= self.comma_seconds or len(self.buf) >= self.comma_chars:
            pos = max(self.buf.rfind(sep) for sep in (", ", "; ", ": "))
            if pos >= 0 and pos + 2 >= self.min_chars:
                return pos + 2
        # 太長了：在逗號或空白處切
        if len(self.buf) > self.max_chars:
            window = self.buf[: self.max_chars]
            for sep in (", ", "; ", ": ", " "):
                pos = window.rfind(sep)
                if pos > self.max_chars // 2:
                    return pos + len(sep)
            return self.max_chars
        return None

    def flush_if_idle(self, now: Optional[float] = None) -> Optional[str]:
        now = time.monotonic() if now is None else now
        if self.buf.strip() and now - self.last_update >= self.idle_seconds:
            return self.flush()
        return None

    def flush(self) -> Optional[str]:
        piece, self.buf = self.buf.strip(), ""
        return piece or None

    @property
    def pending(self) -> str:
        return self.buf


# ---------------------------------------------------------------------------
# 2. 資料結構
# ---------------------------------------------------------------------------
@dataclass
class Segment:
    idx: int
    elapsed: float            # 開始錄音後幾秒
    source: str               # 原文
    translation: Optional[str] = None   # None = 翻譯中
    error: Optional[str] = None


def fmt_time(seconds: float) -> str:
    s = int(seconds)
    return f"{s // 3600:02d}:{s % 3600 // 60:02d}:{s % 60:02d}"


def _message_text(resp) -> str:
    """從 chat.complete 的回應取出文字（content 可能是字串或多段）。"""
    content = resp.choices[0].message.content
    if isinstance(content, str):
        return content.strip()
    parts = []
    for c in content or []:
        t = getattr(c, "text", None)
        if t is None and isinstance(c, dict):
            t = c.get("text")
        if t:
            parts.append(t)
    return "".join(parts).strip()


# ---------------------------------------------------------------------------
# 3. 音訊來源
# ---------------------------------------------------------------------------
def load_wav_as_pcm16(path: str | Path) -> bytes:
    """讀 WAV 檔，轉成 16 kHz、單聲道、16-bit PCM（模擬上課用）。"""
    with wave.open(str(path), "rb") as w:
        n_ch, width, rate, n = w.getnchannels(), w.getsampwidth(), w.getframerate(), w.getnframes()
        raw = w.readframes(n)
    if width != 2:
        raise ValueError("Only 16-bit PCM WAV files are supported")
    data = np.frombuffer(raw, dtype=np.int16).astype(np.float32)
    if n_ch > 1:
        data = data.reshape(-1, n_ch).mean(axis=1)
    if rate != SAMPLE_RATE:
        new_len = int(len(data) * SAMPLE_RATE / rate)
        data = np.interp(np.linspace(0, len(data) - 1, new_len), np.arange(len(data)), data)
    return np.clip(data, -32768, 32767).astype(np.int16).tobytes()


def list_input_devices() -> list[tuple[int, str]]:
    import sounddevice as sd
    return [(i, d["name"]) for i, d in enumerate(sd.query_devices()) if d["max_input_channels"] > 0]


# ---------------------------------------------------------------------------
# 4. 即時翻譯 Session（在背景執行緒跑自己的 asyncio 迴圈）
# ---------------------------------------------------------------------------
TRANSLATE_SYSTEM = (
    "你是大學課堂的即時口譯員。把使用者給的課堂逐字稿片段（英文或法文）翻譯成自然、準確的繁體中文（台灣用語）。\n"
    "規則：\n"
    "1. 片段裡的每一句都要翻譯，包括開場、轉折的短句（例如 Let's start with a simple question.、Now, compare two models.），"
    "不可以省略、摘要或合併句子。不要加解釋、不要重複原文。\n"
    "2. 專有名詞、人名、理論名稱、機構名稱在翻譯中保留原文在括號內，例如：轉換型領導（transformational leadership）。"
    "只使用學術界通用的中文譯名；不確定有沒有通用譯名時，直接保留原文，不要自己創造譯名，也不要用成語或字面意思硬翻。\n"
    "3. 逐字稿是語音辨識結果，可能有錯字或斷句不完整；照字面合理翻譯，不要自行補充原文沒有的內容。\n"
    "4. 「前文」只是幫助你理解上下文：不要翻譯前文，也不要把前文的內容加進翻譯裡。\n"
    "5. 用 JSON 回覆，格式：{\"translation\": \"翻譯\", \"terms\": [{\"en\": \"原文\", \"zh\": \"你用的中文譯名\"}]}。"
    "terms 列出這個片段（不含前文）裡你附了原文的每一個名詞；zh 必須和 translation 裡用的譯名一致，保留原文沒翻的就填原文。沒有就給空陣列。"
)


def glossary_prompt(glossary: dict[str, str]) -> str:
    """把指定譯名加到系統提示後面。"""
    if not glossary:
        return ""
    lines = "\n".join(f"{en} = {zh}" for en, zh in glossary.items())
    return ("\n\n確認過的譯名（只是譯名對照，不是課堂內容）：原文是同樣意思時，一定要使用這裡的中文譯名；"
            "只有在原文明顯是另一個意思時（例如同一個字當一般用語），才依前後文翻譯。\n" + lines)


def background_prompt(brief: str) -> str:
    """把從簡報整理出的課程背景加到系統提示後面。"""
    if not brief.strip():
        return ""
    return (f"\n\n這堂課的背景（從上課簡報整理，只用來判斷領域和用詞，不是課堂內容）：{brief.strip()}\n"
            "用詞請依照這個領域的習慣；專有名詞的譯名以「確認過的譯名」為準。")


def merge_terms(suggested: dict[str, str], fixed: dict[str, str]) -> dict[str, str]:
    """合併簡報譯名和使用者修正的譯名；原文不分大小寫，使用者修正的優先。"""
    merged = {en.lower(): (en, zh) for en, zh in suggested.items()}
    merged.update({en.lower(): (en, zh) for en, zh in fixed.items()})
    return dict(merged.values())


def fix_brief(brief: str, suggested: dict[str, str], fixed: dict[str, str]) -> str:
    """課程背景裡如果用了被使用者修正掉的舊譯名，換成新譯名。"""
    old = {en.lower(): zh for en, zh in suggested.items()}
    for en, zh in fixed.items():
        if old.get(en.lower()) and old[en.lower()] != zh:
            brief = brief.replace(old[en.lower()], zh)
    return brief


def terms_in(text: str, terms: dict[str, str]) -> dict[str, str]:
    """只留原文有出現在 text 裡的名詞。"""
    low = text.lower()
    return {en: zh for en, zh in terms.items() if en.lower() in low}


LANGUAGES = {"English": "英文", "French": "法文"}


def language_prompt(language: str) -> str:
    """告訴模型這堂課主要用哪個語言授課（語音辨識模型沒有語言設定，會自動判斷，所以只用在翻譯和筆記）。"""
    zh = LANGUAGES.get(language)
    if not zh:
        return ""
    return f"\n\n這堂課主要用{zh}授課，逐字稿大多是{zh}（偶爾可能夾雜其他語言），請照{zh}的語意理解後翻譯。"


def build_translate_messages(context: str, text: str, glossary: Optional[dict[str, str]] = None,
                             background: str = "", language: str = "") -> list[dict]:
    user = (f"前文（不用翻譯）：{context}\n\n" if context else "") + f"要翻譯的片段：\n{text}"
    system = TRANSLATE_SYSTEM + language_prompt(language) + background_prompt(background) + glossary_prompt(glossary or {})
    return [{"role": "system", "content": system}, {"role": "user", "content": user}]


def parse_translation(raw: str) -> tuple[str, list[tuple[str, str]]]:
    """解析翻譯模型的 JSON 回覆；模型沒照格式回時，整段當翻譯、沒有名詞。"""
    try:
        data = json.loads(raw)
        text = str(data["translation"]).strip()
        terms = [(str(t["en"]).strip(), str(t["zh"]).strip()) for t in data.get("terms") or []
                 if isinstance(t, dict) and t.get("en") and t.get("zh")]
        return text, terms
    except (ValueError, KeyError, TypeError):
        return raw.strip(), []


class TermBook:
    """
    專有名詞表：翻譯時自動收集「原文 → 譯名」。
    使用者修正過的譯名（fixed）會強制套用到之後的翻譯和筆記；
    沒修正過的也會在之後的翻譯沿用，讓同一個名詞前後譯名一致。
    """

    def __init__(self, fixed: Optional[dict[str, str]] = None, suggested: Optional[dict[str, str]] = None):
        """fixed：使用者修正過的譯名；suggested：從簡報找到的譯名（使用者修正的優先）。"""
        self.lock = threading.Lock()
        self._terms: dict[str, dict] = {}       # key = 原文小寫
        for en, zh in (suggested or {}).items():
            self.add_seen(en, zh)
        for en, zh in (fixed or {}).items():
            self.fix(en, zh)

    def add_seen(self, en: str, zh: str):
        with self.lock:
            self._terms.setdefault(en.lower(), {"en": en, "zh": zh, "fixed": False})

    def fix(self, en: str, zh: str):
        en, zh = en.strip(), zh.strip()
        if en and zh:
            with self.lock:
                self._terms[en.lower()] = {"en": en, "zh": zh, "fixed": True}

    def fixed(self) -> dict[str, str]:
        with self.lock:
            return {t["en"]: t["zh"] for t in self._terms.values() if t["fixed"]}

    def rows(self) -> list[dict]:
        with self.lock:
            return [dict(t) for t in self._terms.values()]

    def relevant(self, text: str) -> dict[str, str]:
        """只挑原文有出現在 text 裡的名詞，避免提示越來越長。"""
        low = text.lower()
        with self.lock:
            return {t["en"]: t["zh"] for k, t in self._terms.items() if k in low}


@dataclass
class SessionConfig:
    api_key: str
    out_dir: Path
    course: str = ""
    stt_model: str = DEFAULT_STT_MODEL
    translate_model: str = DEFAULT_TRANSLATE_MODEL
    target_delay_ms: Optional[int] = None
    input_device: Optional[int] = None       # None = 系統預設麥克風
    wav_path: Optional[str] = None           # 有值 = 用音訊檔模擬
    wav_speed: float = 1.0                   # 模擬播放速度
    min_chars: int = 80
    fixed_terms: dict = field(default_factory=dict)   # 使用者修正過的譯名 {原文: 中文}
    slide_brief: str = ""                              # 從簡報整理的課程背景
    slide_terms: dict = field(default_factory=dict)   # 從簡報找到的譯名 {原文: 中文}
    language: str = "English"                         # 主要授課語言：English / French
    provider: str = "mistral"                         # 用哪家的 API：mistral / gemini（預設 Mistral，研究基準）
    save_audio: bool = True                  # 同時把錄音存成 FLAC（無損壓縮）


CONN_LOG = "connection_log.csv"                  # 每堂課一份（在課程資料夾裡）
CONN_LOG_ALL = "connection_log_all.csv"         # 所有課程一份總表（在 records/ 裡）
CONN_LOG_FIELDS = ["time", "t_seconds", "code_version", "course", "class_folder", "conn_id", "event", "category",
                   "error_type", "attempt", "since_last_text_seconds", "since_lost_seconds", "connect_seconds",
                   "downtime_seconds", "detail",
                   # 以下是研究用的音訊追蹤欄位（加在最後，舊欄位位置不變）
                   "loud_chunks_since_text", "audio_captured_seconds", "audio_handed_seconds",
                   "captured_seq", "dequeued_seq", "handed_seq", "send_returned_seq",
                   "conn_first_handed_seq", "conn_last_handed_seq", "conn_last_send_returned_seq",
                   "queued_chunks", "orphaned_chunks"]
CONFIG_FILE = "config.json"                  # 每堂課一份：程式版本與所有門檻、模型設定（之後比較不同課堂用）

# 研究用量測（量字幕延遲、卡死偵測延遲；之後可以拿 recording.flac 做 trace-driven 模擬）
LATENCY_LOG = "subtitle_latency.csv"               # 每收到一段辨識文字記一筆
LATENCY_LOG_FIELDS = ["wall_time", "t_seconds", "conn_id", "audio_sent_seconds", "audio_captured_seconds",
                      "backlog_seconds", "chars", "text", "handed_seq", "captured_seq"]
TRANSLATION_LOG = "translation_latency.csv"           # 每句話翻譯完（或放棄）記一筆
TRANSLATION_LOG_FIELDS = ["wall_time", "idx", "segment_t_seconds", "start_t_seconds", "done_t_seconds",
                          "queue_seconds", "translate_seconds", "total_seconds", "attempts", "ok",
                          "audio_sent_seconds", "chars", "error"]
HEALTH_LOG = "health_metrics.csv"                # 每秒記一筆
HEALTH_LOG_FIELDS = ["wall_time", "t_seconds", "conn_id", "loud_chunks", "chunks", "text_chars",
                     "since_last_text_seconds", "down", "audio_sent_seconds", "backlog_seconds",
                     "input_overflows", "captured_seq", "dequeued_seq", "handed_seq", "send_returned_seq",
                     "orphaned_chunks"]
ACTIVITY_LOG = "activity_events.csv"      # 上課時手動標記的課堂情境（例如播影片開始／結束），只記錄、不影響偵測
ACTIVITY_LOG_FIELDS = ["wall_time", "t_seconds", "class_folder", "conn_id", "activity", "event", "source",
                       "claimed_offset_seconds", "est_t_seconds", "captured_seq", "handed_seq",
                       "audio_captured_seconds", "note"]
# 卡死偵測與重連策略的版本：偵測邏輯有任何改變就要換這個字串（只加紀錄功能時不用換）
DETECTOR_VERSION = ("stall-v1: reconnect when no text > 30 s and >= 150 chunks (100 ms) with level >= 0.5 "
                    "since last text/connect; backoff min(2*failures, 10) s; server close -> reconnect at once")
ASR_EVENT_LOG = "asr_events.csv"              # 每個不是 text.delta 的 ASR 事件記一筆（text.delta 只記在 subtitle_latency.csv）
ASR_EVENT_LOG_FIELDS = ["wall_time", "t_seconds", "class_folder", "conn_id", "event_type", "text_chars",
                        "seg_start_raw", "seg_end_raw", "handed_seq", "audio_handed_seconds", "detail"]
# 音訊序號（seq）：每段 100 毫秒的音訊在收音時編號，從 0 開始遞增
#   captured  收音並放進 _audio_q 的最後一段
#   dequeued  從 _audio_q 取出的最後一段（在背景執行緒）
#   handed    交給 SDK（Python generator yield 出去）的最後一段；audio_sent_seconds / audio_handed_seconds 就是這個
#   send_returned  SDK 的 send_audio() 已經返回（寫進本機 WebSocket）的最後一段；伺服器有沒有收到：無法觀測（Unknown）
#   orphaned  已經從 _audio_q 取出、但因為連線取消而沒交給 SDK 的段數（取消時的競態，這些音訊不會被辨識）


class _CsvBuffer:
    """先放在記憶體，每幾秒整批寫進 CSV（寫不進去就丟掉，不影響錄音）。"""

    def __init__(self, path: Path, fields: list[str], on_error: Optional[Callable[[str], None]] = None):
        self.path, self.fields = path, fields
        self.rows: list[dict] = []
        self.on_error = on_error
        self.failures = 0                  # 寫檔失敗幾次（失敗的那批資料就丟掉了）

    def add(self, row: dict):
        self.rows.append(row)

    def flush(self):
        rows, self.rows = self.rows, []
        if not rows:
            return
        try:
            append_csv(self.path, self.fields, rows)
        except (OSError, ValueError) as e:
            self.failures += 1
            if self.on_error and self.failures in (1, 10, 100):   # 不要每 5 秒都跳一次
                self.on_error(f"Research log {self.path.name} could not be written "
                              f"({self.failures} time(s), {len(rows)} rows lost): {e}")


def append_csv(path: Path, fields: list[str], rows: list[dict]):
    """把幾筆資料接到 CSV 後面。舊檔的欄位跟現在不一樣，就把舊檔改名保留、開新檔，不會混在一起。"""
    new = not path.exists()
    if not new:
        with open(path, newline="", encoding="utf-8") as f:
            header = next(csv.reader(f), [])
        if header != fields:
            path.rename(path.with_name(f"{path.stem}_old_format_{datetime.now():%Y%m%d-%H%M%S}{path.suffix}"))
            new = True
    with open(path, "a", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        if new:
            w.writeheader()
        w.writerows(rows)


class SilentStall(ConnectionError):
    """連線沒斷也沒報錯，但講話時一直沒有文字回來。"""


class ServerErrorEvent(ConnectionError):
    """伺服器在串流裡送來錯誤事件（附錯誤代碼）。"""

    def __init__(self, message: str, code=None):
        super().__init__(message)
        self.code = code


def error_type(e: BaseException) -> str:
    """錯誤的類別名稱＋代碼（例如 ServerErrorEvent:3001、ConnectionClosedError:1012），統計故障原因用。"""
    code = (getattr(getattr(e, "rcvd", None), "code", None)       # websockets 的 close code（例如 1012）
            or getattr(e, "status_code", None) or getattr(e, "code", None))
    name = type(e).__name__
    return f"{name}:{str(code)[:40]}" if code not in (None, "") else name


def code_version() -> str:
    """目前程式的 git commit（有沒 commit 的修改會加上 -dirty）；查不到就回 unknown。"""
    here = Path(__file__).resolve().parent
    try:
        commit = subprocess.run(["git", "rev-parse", "--short", "HEAD"], cwd=here, capture_output=True,
                                text=True, timeout=3).stdout.strip()
        dirty = subprocess.run(["git", "status", "--porcelain", "--untracked-files=no"], cwd=here,
                               capture_output=True, text=True, timeout=3).stdout.strip()
        return (commit + ("-dirty" if dirty else "")) if commit else "unknown"
    except (OSError, subprocess.SubprocessError):
        return "unknown"


def classify_problem(msg: str) -> str:
    """把錯誤訊息歸類成看得懂的原因，方便之後統計是哪一種問題最多。"""
    m = msg.lower()
    for key, label in (("transcriptiondone", "server ended stream (transcription.done)"),
                       ("streamended", "server ended stream (no error)"),
                       ("toomanyrequests", "server busy (TooManyRequests)"),
                       ("queueoverflow", "server busy (QueueOverflow)"),
                       ("enginedead", "server crash (EngineDead)"),
                       ("1012", "server restart (1012)"),
                       ("service restart", "server restart (1012)"),
                       ("internal error", "server internal error"),
                       ("timeout", "server timeout"),
                       ("no text for", "silent stall (no text while speaking)"),
                       ("microphone", "microphone"),
                       ("device", "microphone")):
        if key in m:
            return label
    if any(k in m for k in ("connect", "network", "socket", "closed", "unreachable", "name resolution",
                            "dns", "ssl", "eof", "reset by peer")):
        return "network / connection closed"
    return "other"


def _pkg_version(name: str) -> str:
    try:
        from importlib.metadata import version
        return version(name)
    except Exception:
        return "unknown"


class LiveSession:
    def __init__(self, cfg: SessionConfig, client_factory: Optional[Callable] = None):
        self.cfg = cfg
        self._client_factory = client_factory
        self.lock = threading.Lock()
        self.segments: list[Segment] = []
        self.partial = ""
        self.language: Optional[str] = None
        self.status = "Not started"
        self.errors: list[str] = []
        self.started_at: Optional[float] = None
        self.ended_at: Optional[float] = None
        self.audio_seconds = 0.0
        self.running = False
        self._stop = threading.Event()
        self._thread: Optional[threading.Thread] = None
        self._buffer = SentenceBuffer(min_chars=cfg.min_chars)
        self._audio_q: "thread_queue.Queue[tuple[int, bytes]]" = thread_queue.Queue()   # (seq, PCM)
        cfg.out_dir.mkdir(parents=True, exist_ok=True)
        self.transcript_path = cfg.out_dir / "transcript_bilingual.md"
        self.source_path = cfg.out_dir / "transcript_source.txt"
        self.audio_path = cfg.out_dir / "recording.flac"
        self._audio_file = None
        self._caffeinate: Optional[subprocess.Popen] = None
        self._producer: Optional[threading.Thread] = None
        self._down_since: Optional[float] = None    # 斷線開始的時間（字幕恢復時算斷了多久）
        self._connect_started_at: Optional[float] = None   # 這條連線開始建立的時間（研究用量測）
        self._code_version = "unknown"
        self._drops = 0
        self._downtime = 0.0
        self._last_text_at = time.monotonic()      # 最後一次收到辨識文字的時間（卡死偵測用）
        self._loud_since_text = 0                   # 那之後收到幾格明顯的說話聲
        # 研究用量測：_last_text_real 不會因為重連而重設（卡死偵測用的 _last_text_at 會）
        self._last_text_real = time.monotonic()
        self._t0 = time.monotonic()
        self.audio_captured_seconds = 0.0          # 收到的音訊（= recording.flac 裡的位置）
        self._conn_id = 0
        self._sec_chunks = self._sec_loud = self._sec_chars = 0   # 這一秒的累計（每秒歸零）
        self._meter_lock = threading.Lock()         # 收音執行緒和健康紀錄同時改計數用
        self._latency_log = _CsvBuffer(cfg.out_dir / LATENCY_LOG, LATENCY_LOG_FIELDS, self._log_error)
        self._health_log = _CsvBuffer(cfg.out_dir / HEALTH_LOG, HEALTH_LOG_FIELDS, self._log_error)
        self._translation_log = _CsvBuffer(cfg.out_dir / TRANSLATION_LOG, TRANSLATION_LOG_FIELDS, self._log_error)
        self._asr_event_log = _CsvBuffer(cfg.out_dir / ASR_EVENT_LOG, ASR_EVENT_LOG_FIELDS, self._log_error)
        self._conn_log_failures = 0
        # 音訊序號追蹤（見 ASR_EVENT_LOG 上面的說明）；都用 _meter_lock 保護
        self._next_seq = 0
        self._captured_seq: Optional[int] = None
        self._dequeued_seq: Optional[int] = None
        self._dequeued_count = 0
        self._handed_seq: Optional[int] = None
        self._handed_count = 0
        self._send_returned_seq: Optional[int] = None
        self._send_returned_count = 0
        self._orphaned_seqs: list[int] = []
        self._conn_stats = {"first": None, "last": None, "send_returned": None}   # 目前這條連線
        self.activity_open: dict[str, float] = {}  # 目前標記為進行中的課堂情境 {"video": 標記的 t_seconds}
        self._sec_overflows = 0                     # 這一秒麥克風 input overflow 幾次
        self._overflows_total = 0
        self._seg_formed_at: dict[int, float] = {}   # 句子斷好（送去翻譯）的時間，算翻譯延遲用
        self.levels: deque[float] = deque(maxlen=20)   # 最近 2 秒的音量（每 100 毫秒一格，0–1），畫音量表用
        self._written_upto = 0
        self.terms = TermBook(cfg.fixed_terms, cfg.slide_terms)

    # ----- 對外介面 -----
    def start(self):
        if self.running:
            return
        self.running = True
        self._stop.clear()
        self.started_at = time.time()
        self._t0 = self._last_text_real = time.monotonic()
        with open(self.transcript_path, "w", encoding="utf-8") as f:
            title = self.cfg.course or "上課逐字稿"
            f.write(f"# {title}\n\n錄音開始：{datetime.now():%Y-%m-%d %H:%M}\n\n")
        self.source_path.write_text("", encoding="utf-8")
        if self.cfg.save_audio:
            try:
                import soundfile as sf
                self._audio_file = sf.SoundFile(self.audio_path, "w", samplerate=SAMPLE_RATE, channels=1,
                                                format="FLAC", subtype="PCM_16")
            except Exception as e:
                self._log_error(f"Could not create the audio file; saving the transcript only: {e}")
        self._keep_awake()
        self._code_version = code_version()
        self._write_config()
        self._log_conn("start", detail=f"model={self.cfg.stt_model} language={self.cfg.language} "
                                       f"stall_seconds={self.STALL_SECONDS} detector={DETECTOR_VERSION.split(':')[0]}")
        self._thread = threading.Thread(target=self._thread_main, daemon=True)
        self._thread.start()

    def _keep_awake(self):
        """錄音期間不讓 Mac 睡眠、螢幕也不自動關（-d 螢幕、-i 系統閒置；-w 程式結束就自動解除）。"""
        if sys.platform == "darwin" and shutil.which("caffeinate"):
            try:
                self._caffeinate = subprocess.Popen(["caffeinate", "-d", "-i", "-w", str(os.getpid())])
            except Exception as e:
                self._log_error(f"Couldn't keep the Mac awake; it may sleep during class: {e}")

    def stop(self):
        self._stop.set()
        self._set_status("Stopping… (finishing the last lines)")

    def mark_activity(self, activity: str, start: bool, seconds_ago: float = 0.0, source: str = "button",
                      note: str = "") -> bool:
        """上課時手動標記課堂情境（例如 activity="video"），寫一筆到 activity_events.csv。
        只是記錄：不會暫停或改變卡死偵測、重連。seconds_ago 是「其實是幾秒前開始／結束的」（漏按時補記）。
        回傳是否寫入成功。"""
        now = time.monotonic()
        t = now - self._t0
        offset = max(float(seconds_ago or 0), 0.0)
        with self.lock:
            if start:
                self.activity_open[activity] = t - offset
            else:
                self.activity_open.pop(activity, None)
        row = {"wall_time": datetime.now().isoformat(timespec="milliseconds"), "t_seconds": round(t, 3),
               "class_folder": self.cfg.out_dir.name, "conn_id": self._conn_id, "activity": activity,
               "event": "start" if start else "end", "source": source,
               "claimed_offset_seconds": round(offset, 1), "est_t_seconds": round(t - offset, 3),
               "note": re.sub(r"\s+", " ", note)[:300]}
        c = self._audio_counters()
        row.update({k: c[k] for k in ("captured_seq", "handed_seq", "audio_captured_seconds")})
        try:
            with self.lock:
                append_csv(self.cfg.out_dir / ACTIVITY_LOG, ACTIVITY_LOG_FIELDS, [row])
            return True
        except (OSError, ValueError) as e:
            self._log_error(f"Could not write {ACTIVITY_LOG}: {e}")
            return False

    def join(self, timeout: Optional[float] = None):
        if self._thread:
            self._thread.join(timeout)

    def elapsed(self) -> float:
        if not self.started_at:
            return 0.0
        return (self.ended_at or time.time()) - self.started_at

    def est_cost(self) -> float:
        return self.audio_seconds / 60 * REALTIME_PRICE_PER_MIN

    def snapshot(self):
        with self.lock:
            return (list(self.segments), self._buffer.pending, self.status, list(self.errors), self.language)

    def full_source_text(self) -> str:
        with self.lock:
            return "\n".join(f"[{fmt_time(s.elapsed)}] {s.source}" for s in self.segments)

    # ----- 內部 -----
    def _set_status(self, s: str):
        with self.lock:
            self.status = s

    def _log_error(self, msg: str):
        with self.lock:
            self.errors.append(f"{datetime.now():%H:%M:%S} {msg}")

    def _log_conn(self, event: str, category: str = "", attempt: Optional[int] = None,
                  downtime: Optional[float] = None, detail: str = "", error: str = "",
                  since_last_text: Optional[float] = None, since_lost: Optional[float] = None,
                  connect: Optional[float] = None, loud: Optional[int] = None):
        """寫一筆連線紀錄到這堂課的 connection_log.csv 和 records/connection_log_all.csv（寫不進去也不影響錄音）。
        t_seconds 和其他研究用 CSV 用同一個起點，可以直接對齊；秒數欄位都用 monotonic 時鐘算。"""
        def sec(x):
            return "" if x is None else round(x, 3)
        row = {"time": datetime.now().isoformat(timespec="milliseconds"),
               "t_seconds": round(time.monotonic() - self._t0, 3), "code_version": self._code_version,
               "course": self.cfg.course, "class_folder": self.cfg.out_dir.name, "conn_id": self._conn_id,
               "event": event, "category": category, "error_type": error,
               "attempt": "" if attempt is None else attempt,
               "since_last_text_seconds": sec(since_last_text), "since_lost_seconds": sec(since_lost),
               "connect_seconds": sec(connect), "downtime_seconds": sec(downtime),
               "detail": re.sub(r"\s+", " ", detail)[:500],
               "loud_chunks_since_text": "" if loud is None else loud}
        row.update(self._audio_counters())
        for path in (self.cfg.out_dir / CONN_LOG, self.cfg.out_dir.parent / CONN_LOG_ALL):
            try:
                with self.lock:
                    append_csv(path, CONN_LOG_FIELDS, [row])
            except (OSError, ValueError) as e:
                self._conn_log_failures += 1
                if self._conn_log_failures in (1, 10, 100):
                    self._log_error(f"Connection log {path.name} could not be written: {e}")

    def _audio_counters(self) -> dict:
        """目前的音訊序號與佇列狀態（研究用；欄位意思見 ASR_EVENT_LOG 上面的說明）。"""
        def v(x):
            return "" if x is None else x
        with self._meter_lock:
            c = self._conn_stats
            return {"audio_captured_seconds": round(self.audio_captured_seconds, 2),
                    "audio_handed_seconds": round(self.audio_seconds, 2),
                    "captured_seq": v(self._captured_seq), "dequeued_seq": v(self._dequeued_seq),
                    "handed_seq": v(self._handed_seq), "send_returned_seq": v(self._send_returned_seq),
                    "conn_first_handed_seq": v(c["first"]), "conn_last_handed_seq": v(c["last"]),
                    "conn_last_send_returned_seq": v(c["send_returned"]),
                    "queued_chunks": self._next_seq - self._dequeued_count,
                    "orphaned_chunks": len(self._orphaned_seqs)}

    def _went_down(self, msg: str, attempt: Optional[int], err: str = ""):
        """斷線或卡死：記一筆，並開始計算斷了多久。"""
        if self._down_since is None:
            self._down_since = time.monotonic()
            self._drops += 1
        # downtime_seconds 只從宣告斷線起算；since_last_text 是偵測時已經多久沒字幕（例如 silent stall ≈ 30 秒）
        self._log_conn("lost", classify_problem(f"{err} {msg}"), attempt, detail=msg, error=err,
                       since_last_text=time.monotonic() - self._last_text_real, loud=self._loud_since_text)

    def _came_back(self):
        """又收到文字（或麥克風恢復）：記一筆恢復，附上斷了多久。"""
        if self._down_since is not None:
            now = time.monotonic()
            down = now - self._down_since
            self._downtime += down
            self._down_since = None
            self._log_conn("recovered", downtime=down, since_lost=down,
                           since_last_text=now - self._last_text_real)

    def _reconnect_started(self, attempt: int):
        """斷線中開始建立新連線：記一筆，附上從宣告斷線起過了多久（含退避等待）。"""
        self._connect_started_at = time.monotonic()
        if self._down_since is not None:
            self._log_conn("reconnect_start", attempt=attempt,
                           since_lost=self._connect_started_at - self._down_since)

    def _session_created(self):
        """伺服器確認新連線建立：記一筆，把「連上」和「第一筆字幕出現（recovered）」分開。"""
        if self._down_since is not None and self._connect_started_at is not None:
            now = time.monotonic()
            self._log_conn("session_created", since_lost=now - self._down_since,
                           connect=now - self._connect_started_at)
        self._connect_started_at = None

    def _write_config(self):
        """把這堂課用的程式版本、模型、所有門檻存成 config.json，之後比較不同課堂時知道條件是否一樣。"""
        try:
            from importlib.metadata import version
            sdk = version("mistralai")
        except Exception:
            sdk = "unknown"
        b = self._buffer
        cfg = {k: v for k, v in vars(self.cfg).items()
               if k not in ("api_key", "out_dir", "fixed_terms", "slide_terms", "slide_brief")}
        data = {
            "code_version": self._code_version,
            "detector_version": DETECTOR_VERSION,
            "started_at": datetime.fromtimestamp(self.started_at).isoformat(timespec="milliseconds"),
            "t_seconds_origin": "t_seconds = 0 at started_at (monotonic clock)",
            "python": sys.version.split()[0], "mistralai_sdk": sdk, "platform": sys.platform,
            "provider": self.cfg.provider,
            "session_config": cfg,
            "glossary": {"fixed_terms": len(self.cfg.fixed_terms), "slide_terms": len(self.cfg.slide_terms),
                         "slide_brief_chars": len(self.cfg.slide_brief)},
            "audio": {"sample_rate": SAMPLE_RATE, "chunk_ms": CHUNK_MS},
            "stall_detection": {"stall_seconds": self.STALL_SECONDS, "loud_chunks": self.STALL_LOUD_CHUNKS,
                                "loud_level": self.LOUD_LEVEL},
            "reconnect_backoff": "wait min(2 * failures, 10) s after each failed STT connection",
            "microphone_reconnect": "no audio for 3 s -> reopen; wait min(failures, 5) s; "
                                    "fall back to default device after 3 failures",
            "sentence_buffer": {"min_chars": b.min_chars, "max_chars": b.max_chars,
                                "idle_seconds": b.idle_seconds, "comma_seconds": b.comma_seconds,
                                "comma_chars": b.comma_chars},
            "translation_retry": "3 attempts; wait 5*(n) s on 429, else 1.5*(n) s",
            "health_log": {"loud_chunks, chunks, text_chars": "counts over the previous 1 s (reset every tick)",
                           "since_last_text_seconds, down, audio_sent_seconds, backlog_seconds":
                               "instantaneous value at the tick",
                           "loud_chunk": f"{CHUNK_MS} ms chunk with level >= {self.LOUD_LEVEL}"},
            "ui_refresh_seconds": 1.0,
            **({"gemini": {"google_genai_sdk": _pkg_version("google-genai"),
                           "live_rotation_seconds": "proactive reconnect every 570 s (Live sessions are capped at 10 min); "
                                                    "each rotation ends with transcription.done -> normal reconnect",
                           "text_events": "only finalized input_transcription (per utterance) is used as text; "
                                          "interim transcripts are ignored",
                           "send_returned": "means session.send_realtime_input returned"}}
               if self.cfg.provider == "gemini" else {}),
            "audio_seq": "each 100 ms chunk gets seq 0,1,2,... at capture; see asr_events.csv / connection_log.csv: "
                         "captured -> dequeued (background thread) -> handed (yielded to SDK) -> send_returned "
                         "(SDK send_audio returned = written to local WebSocket). Server receipt is not observable. "
                         "audio_sent_seconds (old name) = audio_handed_seconds",
        }
        try:
            (self.cfg.out_dir / CONFIG_FILE).write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str),
                                                        encoding="utf-8")
        except OSError as e:
            self._log_error(f"Could not write {CONFIG_FILE}: {e}")

    def _make_client(self):
        if self._client_factory:
            return self._client_factory(self.cfg.api_key)
        if self.cfg.provider == "gemini":
            from gemini_client import GeminiClient       # 介面跟 Mistral client 一樣，下面的流程不用改
            return GeminiClient(self.cfg.api_key, self.cfg.language)
        from mistralai.client import Mistral
        return Mistral(api_key=self.cfg.api_key)

    def _thread_main(self):
        try:
            asyncio.run(self._main())
        except Exception as e:  # 最後防線
            self._log_error(f"Unexpected error: {e!r}")
        finally:
            # 錄音檔由收音執行緒負責關；收音執行緒沒啟動（例如一開始就出錯）才在這裡關
            if self._producer is None and self._audio_file is not None:
                self._audio_file.close()
                self._audio_file = None
            if self._caffeinate is not None:
                self._caffeinate.terminate()
                self._caffeinate = None
            self.running = False
            self.ended_at = time.time()
            self._set_status("Stopped")
            for activity in list(self.activity_open):  # 錄音結束時還標記著進行中 → 自動補一筆結束
                self.mark_activity(activity, start=False, source="session_end",
                                   note="still marked as ongoing when recording stopped")
            if self._down_since is not None:            # 結束時還沒恢復的那次也算進去
                self._downtime += time.monotonic() - self._down_since
                self._down_since = None
            with self._meter_lock:      # 音訊核對：收音 = 佇列剩下 + 交給 SDK + orphaned（＋ 還在背景執行緒手上的）
                audit = (f" captured_chunks={self._next_seq} dequeued_chunks={self._dequeued_count} "
                         f"handed_chunks={self._handed_count} send_returned_chunks={self._send_returned_count} "
                         f"orphaned_chunks={len(self._orphaned_seqs)} orphaned_seqs={self._orphaned_seqs[:50]} "
                         f"input_overflows={self._overflows_total}")
            self._log_conn("end", downtime=self._downtime,
                           detail=f"duration={self.elapsed():.0f}s drops={self._drops} segments={len(self.segments)}"
                                  + audit)

    async def _main(self):
        client = self._make_client()
        self._trans_q: asyncio.Queue = asyncio.Queue()
        producer = self._producer = threading.Thread(target=self._audio_producer, daemon=True)
        producer.start()
        translator = asyncio.create_task(self._translator(client))
        idle = asyncio.create_task(self._idle_flusher())
        stt = asyncio.create_task(self._stt_loop(client))
        health = asyncio.create_task(self._health_logger())
        stop_seen: Optional[float] = None
        try:
            # 按下停止後，最多等 20 秒讓伺服器送回最後的辨識結果
            while not stt.done():
                await asyncio.sleep(0.2)
                if self._stop.is_set():
                    stop_seen = stop_seen or time.monotonic()
                    if time.monotonic() - stop_seen > 20:
                        stt.cancel()
                        break
            try:
                await stt
            except asyncio.CancelledError:
                pass
        finally:
            idle.cancel()
            health.cancel()
            rest = self._buffer.flush()
            if rest:
                self._add_segment(rest)
            await self._trans_q.put(None)       # 通知翻譯結束
            await translator
            self._write_pending()
            # 等收音執行緒把最後的音訊寫進錄音檔、關檔
            await asyncio.get_running_loop().run_in_executor(None, producer.join, 10)
            self._flush_research_logs()

    def _flush_research_logs(self):
        for log in (self._latency_log, self._health_log, self._translation_log, self._asr_event_log):
            log.flush()

    # 音訊：麥克風或檔案 → 錄音檔 ＋ thread queue
    # 錄音檔在這裡寫，所以網路斷掉、還沒送去辨識的音訊也會先存進錄音檔
    def _audio_producer(self):
        try:
            if self.cfg.wav_path:
                pcm = load_wav_as_pcm16(self.cfg.wav_path)
                step = CHUNK_SAMPLES * 2
                for i in range(0, len(pcm), step):
                    if self._stop.is_set():
                        break
                    self._emit_audio(pcm[i:i + step])
                    time.sleep(CHUNK_MS / 1000 / max(self.cfg.wav_speed, 0.1))
                self._stop.set()     # 檔案播完就自動停止
            else:
                self._mic_loop()
        except Exception as e:
            self._log_error(f"Audio input failed: {e}")
            self._stop.set()
        finally:
            if self._audio_file is not None:
                self._audio_file.close()
                self._audio_file = None

    def _emit_audio(self, chunk: bytes):
        with self._meter_lock:                      # 收音順序編號（只在收音執行緒呼叫，但計數會被別的執行緒讀）
            seq = self._next_seq
            self._next_seq += 1
        self._save_audio(chunk)
        self._audio_q.put((seq, chunk))
        # 音量換成 dB，-60 dB（幾乎沒聲音）到 -10 dB（很大聲）對應到 0–1
        rms = float(np.sqrt(np.mean(np.frombuffer(chunk, dtype=np.int16).astype(np.float32) ** 2))) if chunk else 0.0
        db = 20 * np.log10(max(rms, 1.0) / 32768)
        level = min(max((db + 60) / 50, 0.0), 1.0)
        self.levels.append(level)
        loud = level >= self.LOUD_LEVEL
        with self._meter_lock:
            self.audio_captured_seconds += len(chunk) / 2 / SAMPLE_RATE
            self._captured_seq = seq
            self._sec_chunks += 1
            self._sec_loud += loud
        if loud:
            self._loud_since_text += 1

    def _mic_loop(self):
        """
        收麥克風，出問題就自動重新接上，直到按停止為止：
        - 開不了、或錄到一半出錯 → 等一下再開
        - 3 秒都沒收到聲音（例如拔掉外接麥克風）→ 當作斷線，重新開
        - 指定的麥克風連續失敗 3 次 → 改用系統預設麥克風
        """
        import sounddevice as sd
        raw_q: "thread_queue.Queue[bytes]" = thread_queue.Queue()
        device, failures = self.cfg.input_device, 0

        def cb(indata, frames, t, status):
            self._mic_callback(raw_q, indata, status)

        while not self._stop.is_set():
            try:
                with sd.RawInputStream(samplerate=SAMPLE_RATE, channels=1, dtype="int16",
                                       blocksize=CHUNK_SAMPLES, device=device, callback=cb) as stream:
                    last_data = time.monotonic()
                    while not self._stop.is_set():
                        try:
                            self._emit_audio(raw_q.get(timeout=0.2))
                        except thread_queue.Empty:
                            if not stream.active or time.monotonic() - last_data > 3:
                                raise RuntimeError("no sound from the microphone for 3 seconds")
                            continue
                        last_data = time.monotonic()
                        if failures:
                            self._log_error("Microphone reconnected")
                            self._came_back()
                            self._set_status("Recording")
                            failures = 0
            except Exception as e:
                if self._stop.is_set():
                    break
                failures += 1
                self._log_error(f"Microphone problem (attempt {failures}): {e}")
                self._went_down(f"microphone: {e}", failures, error_type(e))
                self._set_status("Microphone disconnected — reconnecting…")
                if device is not None and failures >= 3:
                    self._log_error("Switching to the system default microphone")
                    device = None
                self._stop.wait(min(failures, 5))   # 等的時候按停止也能馬上結束
                try:   # 重新整理裝置清單，才找得到重新插上的麥克風
                    sd._terminate()
                    sd._initialize()
                except Exception:
                    pass
        while not raw_q.empty():
            self._emit_audio(raw_q.get_nowait())

    def _mic_callback(self, raw_q, indata, status):
        """sounddevice 的 callback（音訊執行緒）：只放進佇列、數 overflow，不碰磁碟。"""
        raw_q.put(bytes(indata))
        if status and getattr(status, "input_overflow", False):
            with self._meter_lock:
                self._sec_overflows += 1
                self._overflows_total += 1

    def _dequeue(self, holder: dict):
        """在背景執行緒從 _audio_q 取一段（行為和原本的 get(True, 0.2) 一樣），並記下取出的序號。
        如果等待中這條連線已經取消，這段音訊沒人收了：記成 orphaned（取消時的競態）。"""
        seq, chunk = self._audio_q.get(True, 0.2)
        with self._meter_lock:
            self._dequeued_seq = seq
            self._dequeued_count += 1
            if holder["cancelled"]:
                self._orphaned_seqs.append(seq)
            else:
                holder["item"] = (seq, chunk)
        return seq, chunk

    async def _audio_stream(self, stats: dict) -> AsyncIterator[bytes]:
        """把 _audio_q 裡的音訊依序交給 SDK（FIFO，格式不變：只 yield PCM bytes）。stats 是這條連線的序號統計。"""
        loop = asyncio.get_running_loop()
        while True:
            holder = {"cancelled": False, "item": None}
            try:
                seq, chunk = await loop.run_in_executor(None, self._dequeue, holder)
            except thread_queue.Empty:
                if self._stop.is_set():
                    return
                continue
            except BaseException:
                # 連線被取消：背景執行緒可能已經（或稍後才）從 _audio_q 取走一段 → 那段記成 orphaned
                with self._meter_lock:
                    holder["cancelled"] = True
                    if holder["item"] is not None:
                        self._orphaned_seqs.append(holder["item"][0])
                        holder["item"] = None
                raise
            with self._meter_lock:
                holder["item"] = None
                self._handed_seq = seq
                self._handed_count += 1
                if stats["first"] is None:
                    stats["first"] = seq
                stats["last"] = seq
            self.audio_seconds += len(chunk) / 2 / SAMPLE_RATE
            yield chunk
            # SDK 回來要下一段 = 這一段的 send_audio() 已經返回（寫進本機 WebSocket；伺服器有沒有收到無法觀測）
            with self._meter_lock:
                self._send_returned_seq = stats["send_returned"] = seq
                self._send_returned_count += 1

    def _save_audio(self, chunk: bytes):
        """收到的每一段音訊寫進錄音檔（在收音執行緒呼叫，跟網路狀況無關）。"""
        if self._audio_file is None:
            return
        try:
            self._audio_file.write(np.frombuffer(chunk, dtype=np.int16))
        except Exception as e:
            self._log_error(f"Writing the audio file failed; saving the transcript only from now on: {e}")
            self._audio_file = None

    # 卡死偵測：麥克風這段時間一直有人在講話，卻超過 STALL_SECONDS 秒沒有任何文字回來 → 當作連線卡死，重連
    STALL_SECONDS = 30
    STALL_LOUD_CHUNKS = 150        # 15 秒（每格 100 毫秒）音量明顯的聲音
    LOUD_LEVEL = 0.5               # 約 -35 dB，一般講話的音量；教室底噪通常在這以下

    async def _stt_loop(self, client):
        from mistralai.client.models import (AudioFormat, RealtimeTranscriptionError,
                                             TranscriptionStreamDone, TranscriptionStreamLanguage,
                                             TranscriptionStreamSegmentDelta, TranscriptionStreamTextDelta)
        fmt = AudioFormat(encoding="pcm_s16le", sample_rate=SAMPLE_RATE)
        failures = 0

        async def consume(stream):
            nonlocal failures
            async for ev in stream:
                if isinstance(ev, TranscriptionStreamTextDelta):
                    failures = 0
                    self._came_back()
                    self._record_text(ev.text)
                    self._last_text_at = time.monotonic()
                    self._loud_since_text = 0
                    if not self._stop.is_set():
                        self._set_status("Recording")
                    with self.lock:
                        pieces = self._buffer.add(ev.text)
                    for p in pieces:
                        self._add_segment(p)
                    continue
                self._record_asr_event(ev)          # 研究用：其他事件都記進 asr_events.csv（行為不受影響）
                if isinstance(ev, TranscriptionStreamLanguage):
                    with self.lock:
                        self.language = ev.audio_language
                elif isinstance(ev, TranscriptionStreamDone):
                    return "TranscriptionDone"
                elif isinstance(ev, RealtimeTranscriptionError):
                    # 伺服器那邊出錯（例如逾時、辨識引擎當掉）：這條連線通常已經壞了，
                    # 直接斷開重連，不要一直等；還沒送出的音訊會在重連後補送
                    msg = ev.error.message
                    raise ServerErrorEvent(f"speech recognition server error: {getattr(msg, 'detail', msg)}",
                                           code=getattr(ev.error, "code", None))
                elif getattr(ev, "type", None) in ("session.created", "session.updated"):
                    if ev.type == "session.created":
                        self._session_created()
                    if not self._stop.is_set():
                        self._set_status("Recording")
                # TranscriptionStreamSegmentDelta 和未知事件：原本就不處理，現在只多記一筆
            return "StreamEnded"                    # 伺服器結束了連線（沒有送 transcription.done）

        while True:
            self._set_status("Connecting…")
            try:
                stats = {"first": None, "last": None, "send_returned": None}   # 這條連線實際拿到的音訊序號
                with self._meter_lock:
                    self._conn_stats = stats
                kwargs = dict(audio_stream=self._audio_stream(stats), model=self.cfg.stt_model, audio_format=fmt)
                if self.cfg.target_delay_ms:
                    kwargs["target_streaming_delay_ms"] = self.cfg.target_delay_ms
                self._last_text_at = time.monotonic()
                self._loud_since_text = 0
                self._conn_id += 1
                self._reconnect_started(failures)
                consumer = asyncio.create_task(consume(client.audio.realtime.transcribe_stream(**kwargs)))
                # 一邊收結果，一邊每秒檢查有沒有卡死（連線沒斷、也沒報錯，但一直不回文字）
                while not consumer.done():
                    await asyncio.sleep(1)
                    stalled = (not self._stop.is_set()
                               and time.monotonic() - self._last_text_at > self.STALL_SECONDS
                               and self._loud_since_text >= self.STALL_LOUD_CHUNKS)
                    if stalled:
                        # 研究用：判定卡死的當下先記一筆（lost 要等 SDK 關完連線才記，中間的時間就是關閉花的時間）
                        self._log_conn("stall_detected", attempt=failures + 1, loud=self._loud_since_text,
                                       since_last_text=time.monotonic() - self._last_text_real,
                                       detail=f"since_connect_or_text_seconds="
                                              f"{time.monotonic() - self._last_text_at:.1f}")
                        consumer.cancel()
                        try:
                            await consumer
                        except (asyncio.CancelledError, Exception):
                            pass
                        raise SilentStall(f"no text for {self.STALL_SECONDS}s although the microphone "
                                              "hears speech — restarting the connection")
                ended = await consumer          # 把 consume 裡的錯誤丟出來
                if self.cfg.provider == "gemini" and ended == "TranscriptionDone":
                    # Gemini Live 一條連線最多 10 分鐘，轉接層每 9.5 分鐘主動收尾換新連線：
                    # 這是預定的換連線，不是故障；另外記成 planned_rotation，不算進斷線（行為不變，照常馬上重連）
                    ended = "PlannedRotation"
                    if not self._stop.is_set():
                        self._log_conn("planned_rotation", category="planned rotation (Gemini Live 10-min limit)",
                                       detail="proactive reconnect before the 10-minute Live session cap")
                elif not self._stop.is_set():
                    # 研究用：伺服器自己結束連線（沒有錯誤）→ 原本就直接重連（不等待）；現在多記一筆 lost
                    self._went_down(f"server ended the stream without error ({ended})", None, ended)
            except Exception as e:
                # 沒按停止就一直重連（例如教室 Wi-Fi 斷了一陣子）；斷線期間的音訊會排隊，連上後一起補送
                failures += 1
                self._log_error(f"Connection lost (attempt {failures}): {e}")
                self._went_down(str(e), failures, error_type(e))
                self._set_status(f"Connection lost — reconnecting (attempt {failures})…")
                for _ in range(min(2 * failures, 10) * 5):
                    if self._stop.is_set():
                        break
                    await asyncio.sleep(0.2)
            if self._stop.is_set():
                return
            # 伺服器結束了這次連線但使用者沒按停止 → 自動重連

    # ----- 研究用量測（subtitle_latency.csv、health_metrics.csv）-----
    def _record_text(self, text: str):
        now = time.monotonic()
        self._last_text_real = now
        self._sec_chars += len(text)
        self._latency_log.add({
            "wall_time": datetime.now().isoformat(timespec="milliseconds"),
            "t_seconds": round(now - self._t0, 3), "conn_id": self._conn_id,
            "audio_sent_seconds": round(self.audio_seconds, 2),
            "audio_captured_seconds": round(self.audio_captured_seconds, 2),
            "backlog_seconds": round(self.audio_captured_seconds - self.audio_seconds, 2),
            "chars": len(text), "text": text,
            "handed_seq": "" if self._handed_seq is None else self._handed_seq,
            "captured_seq": "" if self._captured_seq is None else self._captured_seq})

    def _record_asr_event(self, ev):
        """不是 text.delta 的 ASR 事件記一筆到 asr_events.csv（text.delta 只在 subtitle_latency.csv，不會重複）。
        segment 的 start/end 原樣保留：還不知道它是以這條連線開頭、還是整段音訊為 0。"""
        now = time.monotonic()
        etype = getattr(ev, "type", None) or type(ev).__name__
        text = getattr(ev, "text", None)
        detail = ""
        session = getattr(ev, "session", None)
        if session is not None:
            detail = (f"request_id={getattr(session, 'request_id', '')} model={getattr(session, 'model', '')} "
                      f"target_streaming_delay_ms={getattr(session, 'target_streaming_delay_ms', '')}")
        elif getattr(ev, "audio_language", None):
            detail = f"audio_language={ev.audio_language}"
        elif getattr(ev, "error", None) is not None and hasattr(ev.error, "code"):
            msg = ev.error.message
            detail = f"code={ev.error.code} message={getattr(msg, 'detail', msg)}"
        elif hasattr(ev, "content"):                # UnknownRealtimeEvent
            detail = f"error={getattr(ev, 'error', '')} content={str(ev.content)[:300]}"
        elif isinstance(text, str):
            detail = text if not hasattr(ev, "segments") else f"segments={len(ev.segments or [])}"

        def raw(x):
            return "" if x is None or type(x).__name__ in ("Unset", "UnsetType") else x
        self._asr_event_log.add({
            "wall_time": datetime.now().isoformat(timespec="milliseconds"),
            "t_seconds": round(now - self._t0, 3), "class_folder": self.cfg.out_dir.name,
            "conn_id": self._conn_id, "event_type": etype,
            "text_chars": len(text) if isinstance(text, str) else "",
            "seg_start_raw": raw(getattr(ev, "start", None)), "seg_end_raw": raw(getattr(ev, "end", None)),
            "handed_seq": "" if self._handed_seq is None else self._handed_seq,
            "audio_handed_seconds": round(self.audio_seconds, 2),
            "detail": re.sub(r"\s+", " ", str(detail))[:500]})

    def _record_translation(self, seg: Segment, started: float, attempts: int):
        """一句話翻譯完（或三次都失敗）：記下斷句時間、排隊多久、翻譯花多久、重試幾次。"""
        now = time.monotonic()
        formed = self._seg_formed_at.pop(seg.idx, started)
        self._translation_log.add({
            "wall_time": datetime.now().isoformat(timespec="milliseconds"), "idx": seg.idx,
            "segment_t_seconds": round(formed - self._t0, 3), "start_t_seconds": round(started - self._t0, 3),
            "done_t_seconds": round(now - self._t0, 3), "queue_seconds": round(started - formed, 3),
            "translate_seconds": round(now - started, 3), "total_seconds": round(now - formed, 3),
            "attempts": attempts, "ok": int(not seg.error), "audio_sent_seconds": round(seg.elapsed, 2),
            "chars": len(seg.source), "error": re.sub(r"\s+", " ", seg.error or "")[:200]})

    async def _health_logger(self):
        """每秒記一筆：這一秒有幾格明顯說話聲、收到幾個字、多久沒字幕、是不是斷線中；每 5 秒寫檔一次。"""
        next_tick = time.monotonic() + 1
        n = 0
        try:
            while True:
                await asyncio.sleep(max(next_tick - time.monotonic(), 0))
                next_tick += 1
                now = time.monotonic()
                with self._meter_lock:
                    chunks, self._sec_chunks = self._sec_chunks, 0
                    loud, self._sec_loud = self._sec_loud, 0
                    overflows, self._sec_overflows = self._sec_overflows, 0
                    captured = self.audio_captured_seconds
                    seqs = {k: "" if v is None else v for k, v in (
                        ("captured_seq", self._captured_seq), ("dequeued_seq", self._dequeued_seq),
                        ("handed_seq", self._handed_seq), ("send_returned_seq", self._send_returned_seq))}
                    orphaned = len(self._orphaned_seqs)
                chars, self._sec_chars = self._sec_chars, 0
                self._health_log.add({
                    "wall_time": datetime.now().isoformat(timespec="milliseconds"),
                    "t_seconds": round(now - self._t0, 3), "conn_id": self._conn_id,
                    "loud_chunks": loud, "chunks": chunks, "text_chars": chars,
                    "since_last_text_seconds": round(now - self._last_text_real, 1),
                    "down": int(self._down_since is not None),
                    "audio_sent_seconds": round(self.audio_seconds, 2),
                    "backlog_seconds": round(captured - self.audio_seconds, 2),
                    "input_overflows": overflows, **seqs, "orphaned_chunks": orphaned})
                n += 1
                if n % 5 == 0:
                    self._flush_research_logs()
        except asyncio.CancelledError:
            pass

    async def _idle_flusher(self):
        while True:
            await asyncio.sleep(0.5)
            with self.lock:
                piece = self._buffer.flush_if_idle()
            if piece:
                self._add_segment(piece)

    def _add_segment(self, text: str):
        with self.lock:
            seg = Segment(idx=len(self.segments), elapsed=self.audio_seconds, source=text)
            self.segments.append(seg)
            self._seg_formed_at[seg.idx] = time.monotonic()
        with open(self.source_path, "a", encoding="utf-8") as f:
            f.write(f"[{fmt_time(seg.elapsed)}] {text}\n")
        self._trans_q.put_nowait(seg)

    async def _translator(self, client):
        while True:
            seg = await self._trans_q.get()
            if seg is None:
                return
            with self.lock:
                context = " ".join(s.source for s in self.segments[max(0, seg.idx - 2):seg.idx])
            started = time.monotonic()
            attempts = 0
            for attempt in range(3):
                attempts = attempt + 1
                try:
                    glossary = self.terms.relevant(f"{context} {seg.source}")
                    brief = fix_brief(self.cfg.slide_brief, self.cfg.slide_terms, self.terms.fixed())
                    resp = await self._client_chat(client, self.cfg.translate_model,
                                                   build_translate_messages(context, seg.source, glossary, brief,
                                                                            self.cfg.language),
                                                   temperature=0.2, response_format={"type": "json_object"})
                    text, terms = parse_translation(_message_text(resp))
                    for en, zh in terms:
                        self.terms.add_seen(en, zh)
                    with self.lock:
                        seg.translation = text
                    break
                except Exception as e:
                    wait = 5 * (attempt + 1) if "429" in str(e) else 1.5 * (attempt + 1)
                    if attempt == 2:
                        with self.lock:
                            seg.error = f"Translation failed: {e}"
                            seg.translation = ""
                        self._log_error(f"Translation failed: {e}")
                    else:
                        await asyncio.sleep(wait)
            self._record_translation(seg, started, attempts)
            self._write_pending()

    @staticmethod
    async def _client_chat(client, model, messages, **kw):
        return await client.chat.complete_async(model=model, messages=messages, **kw)

    def _write_pending(self):
        """依序把已翻譯好的段落寫入雙語逐字稿（避免順序錯亂）。"""
        with self.lock:
            ready = []
            while self._written_upto < len(self.segments) and self.segments[self._written_upto].translation is not None:
                ready.append(self.segments[self._written_upto])
                self._written_upto += 1
        if ready:
            with open(self.transcript_path, "a", encoding="utf-8") as f:
                for s in ready:
                    f.write(f"**[{fmt_time(s.elapsed)}]** {s.source}\n\n> {s.translation or '（翻譯失敗）'}\n\n")


# ---------------------------------------------------------------------------
# 5. 下課後產生筆記
# ---------------------------------------------------------------------------
NOTES_CHUNK_SYSTEM = (
    "你是研究生的上課筆記助理。以下是一段課堂語音辨識逐字稿（英文或法文，含時間戳）。"
    "請用繁體中文（台灣用語）整理這一段的詳細筆記。\n"
    "要求：\n"
    "- 只根據逐字稿內容，不要加入逐字稿沒有的資訊（課程背景和譯名對照只用來決定用詞）；"
    "聽不清楚或不確定的地方標註「（逐字稿不清楚）」。\n"
    "- 依主題分小節，保留老師舉的例子、數字、人名、理論名稱（附原文）。\n"
    "- 老師提到的作業、考試、截止日、課前準備，全部列出並附時間戳。\n"
    "- 最後列出這段出現的專有名詞（原文｜中文｜一句話解釋）。\n"
    "- 專有名詞只用學術界通用的中文譯名；不確定有沒有通用譯名時，中文欄直接寫原文，不要自己創造譯名，也不要用成語或字面意思硬翻。"
)

NOTES_MERGE_SYSTEM = (
    "你是研究生的上課筆記助理。以下是同一堂課依時間順序分段整理的筆記。"
    "請合併成一份完整、詳細的繁體中文（台灣用語）Markdown 筆記。\n"
    "結構：\n"
    "# 課程名稱與日期\n"
    "## 一、本堂課摘要（5–8 句）\n"
    "## 二、詳細筆記（依主題分節，保留例子、數字、理論名稱原文）\n"
    "## 三、作業／考試／課前準備／重要提醒（附時間戳；沒有就寫「本堂課未提到」）\n"
    "## 四、專有名詞表（表格：原文｜中文｜解釋）\n"
    "## 五、逐字稿中不清楚、需要再確認的地方\n"
    "規則：只根據提供的內容，不要自行補充；重複內容要合併，不要遺漏。"
    "課程背景和譯名對照只用來決定用詞，裡面有但逐字稿沒講到的東西不要寫進筆記。\n"
    "同一個專有名詞全篇只用一種譯名；各段譯名不同時，選學術界通用的那個，不確定就保留原文。"
)


def _tidy_markdown(text: str) -> str:
    """把模型輸出中連續重複的分隔線（---）合併成一條。"""
    return re.sub(r"(?:^---[ \t]*\n\s*){2,}", "---\n\n", text, flags=re.M)


def chunk_text(text: str, max_chars: int = 12000) -> list[str]:
    lines, chunks, cur = text.splitlines(), [], ""
    for line in lines:
        if cur and len(cur) + len(line) + 1 > max_chars:
            chunks.append(cur)
            cur = ""
        cur += line + "\n"
    if cur.strip():
        chunks.append(cur)
    return chunks


def generate_notes(api_key: str, transcript: str, course: str, date_str: str,
                   model: str = DEFAULT_NOTES_MODEL, client=None,
                   progress: Optional[Callable[[str], None]] = None,
                   cache_dir: Optional[Path] = None, retries: int = 4,
                   glossary: Optional[dict[str, str]] = None, background: str = "",
                   language: str = "") -> str:
    """
    分段整理逐字稿，再合併成一份筆記。
    cache_dir：每段整理好就先存檔；中途失敗後再按一次，已完成的段落直接沿用，不會重複扣費。
    glossary：要照用的譯名（使用者修正過的、從簡報找到的）。background：從簡報整理的課程背景。
    """
    if client is None:
        from mistralai.client import Mistral
        client = Mistral(api_key=api_key)
    say = progress or (lambda s: None)
    chunks = chunk_text(transcript)
    if not chunks:
        return "(No transcript content)"

    def ask(system, user):
        for attempt in range(retries):
            try:
                resp = client.chat.complete(model=model, temperature=0.2,
                                            messages=[{"role": "system", "content": system},
                                                      {"role": "user", "content": user}])
                return _message_text(resp)
            except Exception as e:
                # 金鑰無效、沒有權限：重試也沒用，直接回報
                if attempt == retries - 1 or "401" in str(e) or "403" in str(e):
                    raise
                wait = 10 * (attempt + 1) if "429" in str(e) else 3 * (attempt + 1)
                say(f"⚠️ Request failed, retrying in {wait}s (attempt {attempt + 1}): {e}")
                time.sleep(wait)

    # 只給逐字稿裡真的出現過的名詞，避免模型把簡報上的名詞當成上課內容寫進筆記
    extra = (language_prompt(language) + background_prompt(background)
             + glossary_prompt(terms_in(transcript, glossary or {})))
    chunk_system = NOTES_CHUNK_SYSTEM + extra
    merge_system = NOTES_MERGE_SYSTEM + extra

    def cache_file(chunk: str) -> Optional[Path]:
        if cache_dir is None:
            return None
        key = hashlib.sha256(f"{model}\n{chunk_system}\n{chunk}".encode("utf-8")).hexdigest()[:16]
        return cache_dir / f"{key}.md"

    if cache_dir is not None:
        cache_dir.mkdir(parents=True, exist_ok=True)
    partials = []
    for i, c in enumerate(chunks, 1):
        f = cache_file(c)
        if f is not None and f.exists():
            say(f"Part {i}/{len(chunks)} was already done, reusing it")
            partials.append(f.read_text(encoding="utf-8"))
            continue
        say(f"Summarizing part {i}/{len(chunks)}…")
        part = ask(chunk_system, c)
        if f is not None:
            f.write_text(part, encoding="utf-8")
        partials.append(part)
    say("Merging into the final notes…")
    joined = "\n\n---\n\n".join(f"【第 {i} 段】\n{p}" for i, p in enumerate(partials, 1))
    return _tidy_markdown(ask(merge_system, f"課程名稱：{course or '（未填）'}\n日期：{date_str}\n\n{joined}"))


# ---------------------------------------------------------------------------
# 6. 上課前讀簡報：整理課程背景和專有名詞
# ---------------------------------------------------------------------------
SLIDES_SYSTEM = (
    "你是大學課程助教。以下是一堂課的簡報文字（英文或法文）。用 JSON 回覆，格式：\n"
    "{\"summary\": \"…\", \"terms\": [{\"en\": \"原文\", \"zh\": \"繁體中文譯名\"}]}\n"
    "summary：用繁體中文（台灣用語）2–3 句說明這堂課的學科領域、主題和重點概念。\n"
    "terms：列出簡報裡的專有名詞、理論、概念、人名、機構名，最多 80 個。規則：\n"
    "- 只列完整的專業名詞（例如 merit goods、arm's length principle），不要列 board、goods 這種一般單字\n"
    "- zh 用台灣學術界通用的譯名；人名、機構名沒有通用譯名就保留原文；不確定的也保留原文，不要自己創造譯名\n"
    "- zh 只填一個譯名，不要加括號說明或「或稱…」"
)
SLIDES_MAX_CHARS = 40000


def extract_slide_text(filename: str, data: bytes) -> str:
    """從 PDF 或 PPTX 簡報取出文字。"""
    import io
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        pages = PdfReader(io.BytesIO(data)).pages
        return "\n\n".join(f"[第 {i} 頁]\n{p.extract_text() or ''}" for i, p in enumerate(pages, 1))
    if name.endswith(".pptx"):
        from pptx import Presentation
        out = []
        for i, slide in enumerate(Presentation(io.BytesIO(data)).slides, 1):
            texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                texts.append("備註：" + slide.notes_slide.notes_text_frame.text)
            out.append(f"[第 {i} 頁]\n" + "\n".join(texts))
        return "\n\n".join(out)
    raise ValueError("Only PDF or PPTX files are supported")


_PAGE_MARK = re.compile(r"\[第 \d+ 頁\]")
SLIDES_MAX_PARTS = 8          # 最多分 8 批（每批約 80–100 頁），避免超大的檔案花太多錢
SLIDES_MERGE_SYSTEM = (
    "以下是同一份簡報分批整理出的摘要。合併成一段繁體中文（台灣用語）2–3 句，"
    "說明這堂課的學科領域、主題和重點概念。只輸出摘要。"
)


def split_slide_text(text: str, max_chars: int = SLIDES_MAX_CHARS) -> list[str]:
    """把簡報文字切成每批不超過 max_chars 的幾批，盡量在頁與頁之間切（單頁太長才從中間切）。"""
    parts: list[str] = []
    cur = ""
    for page in re.split(r"\n\n(?=\[第 \d+ 頁\])", text):
        while len(page) > max_chars:
            if cur:
                parts.append(cur)
                cur = ""
            parts.append(page[:max_chars])
            page = page[max_chars:]
        if cur and len(cur) + 2 + len(page) > max_chars:
            parts.append(cur)
            cur = page
        else:
            cur = f"{cur}\n\n{page}" if cur else page
    if cur.strip():
        parts.append(cur)
    return parts


def analyze_slides(api_key: str, text: str, model: str = DEFAULT_SLIDES_MODEL,
                   client=None) -> tuple[str, dict[str, str], str]:
    """
    讀簡報文字，回傳（課程背景摘要, {原文: 中文譯名}, 提醒訊息）。
    簡報太長就分批同時讀，再把名詞表合併、摘要整合成一段；超過 SLIDES_MAX_PARTS 批才會略過後面的頁面。
    """
    if not text.strip():
        raise ValueError("No text found in the slides (they may be scanned images)")
    parts = split_slide_text(text)
    warning = ""
    if len(parts) > SLIDES_MAX_PARTS:
        total = len(_PAGE_MARK.findall(text))
        parts = parts[:SLIDES_MAX_PARTS]
        read = sum(len(_PAGE_MARK.findall(p)) for p in parts)
        warning = (f"Slides very long: only the first {read} of {total} slides were read. "
                   "Terms on later slides won't be picked up automatically.")
    if client is None:
        from mistralai.client import Mistral
        client = Mistral(api_key=api_key)

    def ask(messages, json_mode: bool) -> str:
        for attempt in range(3):
            try:
                kw = {"response_format": {"type": "json_object"}} if json_mode else {}
                resp = client.chat.complete(model=model, temperature=0.1, messages=messages, **kw)
                return _message_text(resp)
            except Exception as e:
                if attempt == 2 or "401" in str(e) or "403" in str(e):
                    raise
                time.sleep(10 * (attempt + 1) if "429" in str(e) else 3 * (attempt + 1))

    def read_part(part: str) -> dict:
        return json.loads(ask([{"role": "system", "content": SLIDES_SYSTEM},
                               {"role": "user", "content": part}], json_mode=True))

    from concurrent.futures import ThreadPoolExecutor
    with ThreadPoolExecutor(max_workers=4) as pool:
        results = list(pool.map(read_part, parts))

    # 合併各批的名詞：原文不分大小寫，先出現的譯名為準
    terms: dict[str, str] = {}
    seen: set[str] = set()
    for data in results:
        for t in data.get("terms") or []:
            if isinstance(t, dict) and t.get("en") and t.get("zh"):
                en = str(t["en"]).strip()
                # 模型偶爾會寫成「A（或稱「B」）」，只留第一個譯名
                zh = re.sub(r"\s*[（(]\s*(或稱|又稱|或|亦稱).*$", "", str(t["zh"])).strip() or en
                if en.lower() not in seen:
                    seen.add(en.lower())
                    terms[en] = zh
    summaries = [str(d.get("summary", "")).strip() for d in results if d.get("summary")]
    if len(summaries) <= 1:
        summary = summaries[0] if summaries else ""
    else:
        summary = ask([{"role": "system", "content": SLIDES_MERGE_SYSTEM},
                       {"role": "user", "content": "\n\n".join(summaries)}], json_mode=False).strip()
    return summary, terms, warning
